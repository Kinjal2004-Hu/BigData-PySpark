from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

doc = Document("BDA_MiniProject_Report_Kinjal_Ashish.docx")

def iter_blocks(parent):
    for child in parent.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

print("=== FRONT MATTER (first 60 paras) ===")
for i, p in enumerate(doc.paragraphs[:60]):
    txt = p.text.replace('\n', ' / ')
    if txt.strip() or 'graphic' in p._p.xml:
        b = 'B' if any(r.bold for r in p.runs) else ' '
        sizes = {str(r.font.size.pt) if r.font.size else '-' for r in p.runs}
        print(f"P{i:03d}[{b}{sizes}] {txt[:90]}")

print()
print("=== BODY OUTLINE (bold headings + subheads + images) ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    has_img = 'graphic' in p._p.xml or 'w:drawing' in p._p.xml
    if has_img:
        print(f"P{i:03d} <IMAGE>")
    if txt and any(r.bold for r in p.runs):
        size = next((r.font.size.pt for r in p.runs if r.font.size), 12)
        if size >= 14 or txt[0].isdigit() or txt.startswith('5.'):
            print(f"P{i:03d} [H{size:g}] {txt[:70]}")

print()
print("=== IMAGE RELS ===")
imgs = [(rel.rId, rel.target_ref, len(rel.target_part.blob)) for rel in doc.part.rels.values() if 'image' in rel.reltype]
for r in imgs:
    print(r)

print()
print("=== CODE FONT CHECK (sample) ===")
code_paras = [p for p in doc.paragraphs if p.runs and p.runs[0].font.name == 'Consolas']
print(f"Consolas paragraphs: {len(code_paras)}")
if code_paras:
    r = code_paras[5].runs[0]
    print(f"sample: '{r.text[:60]}' size={r.font.size.pt} color={r.font.color.rgb}")

print()
print("=== LAST 12 PARAGRAPHS ===")
for p in doc.paragraphs[-12:]:
    if p.text.strip():
        print(' ', p.text[:90])

print()
print("Total paragraphs:", len(doc.paragraphs))
